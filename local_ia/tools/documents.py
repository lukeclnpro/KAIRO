"""Structured document, table, archive, and local database operations."""

from __future__ import annotations

import csv
import html
import io
import json
import math
import re
import sqlite3
import stat
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

import file_commands

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_BYTES = 200 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5000
MAX_TABLE_ROWS = 20000
MAX_SQL_ROWS = 1000


def _resolve(path, allowed_roots=None):
    return file_commands._resolve(path, allowed_roots)


def _read_bytes(path, allowed_roots=None):
    target = _resolve(path, allowed_roots)
    if not target.is_file():
        raise FileNotFoundError(f"Fichier introuvable : {target}")
    size = target.stat().st_size
    if size > MAX_DOCUMENT_BYTES:
        raise ValueError("Le document dépasse la limite de 20 Mo.")
    return target, target.read_bytes()


def _write_bytes(path, payload, allowed_roots=None):
    target = _resolve(path, allowed_roots)
    if len(payload) > MAX_DOCUMENT_BYTES:
        raise ValueError("Le document produit dépasse la limite de 20 Mo.")
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        with target.open("xb") as output:
            output.write(payload)
    except FileExistsError as error:
        raise FileExistsError(f"Le fichier existe déjà : {target}") from error
    return {"path": str(target), "size": len(payload), "format": target.suffix.lstrip("."), "created": True}


def _text_from_document(path, raw):
    extension = path.suffix.casefold()
    if extension in {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".xml"}:
        return raw.decode("utf-8-sig")
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise ValueError("Lecture PDF indisponible. Installe les dépendances du projet.") from error
        reader = PdfReader(io.BytesIO(raw))
        pages = reader.pages[:200]
        return "\n\n".join(page.extract_text() or "" for page in pages)
    if extension == ".docx":
        try:
            from docx import Document
        except ImportError as error:
            raise ValueError("Lecture DOCX indisponible. Installe les dépendances du projet.") from error
        document = Document(io.BytesIO(raw))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            paragraphs.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("Lecture XLSX indisponible. Installe les dépendances du projet.") from error
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        chunks = []
        for sheet in workbook.worksheets:
            chunks.append(f"## {sheet.title}")
            chunks.extend("\t".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
        workbook.close()
        return "\n".join(chunks)
    if extension == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as error:
            raise ValueError("Lecture PPTX indisponible. Installe les dépendances du projet.") from error
        presentation = Presentation(io.BytesIO(raw))
        slide_text = []
        for slide in presentation.slides:
            shape_text = []
            for shape in slide.shapes:
                text_frame = getattr(shape, "text_frame", None)
                text = getattr(text_frame, "text", "")
                if text:
                    shape_text.append(text)
            slide_text.append("\n".join(shape_text))
        return "\n\n".join(slide_text)
    raise ValueError(f"Format de document non pris en charge : {extension or '(sans extension)'}")


def _ocr_pdf(raw, max_pages=20):
    try:
        import pytesseract
        from pdf2image import convert_from_bytes
        from pypdf import PdfReader
    except ImportError as error:
        raise ValueError("OCR PDF indisponible. Installe les dépendances Python OCR du projet.") from error
    page_count = min(len(PdfReader(io.BytesIO(raw)).pages), max_pages)
    if not page_count:
        return ""
    try:
        images = convert_from_bytes(raw, first_page=1, last_page=page_count, fmt="png", dpi=150)
        return "\n\n".join(pytesseract.image_to_string(image, timeout=30) for image in images)
    except Exception as error:
        raise ValueError("OCR PDF impossible. Vérifie que Tesseract et Poppler sont installés.") from error


def _rows_from_path(path, allowed_roots=None) -> tuple[Path, list[dict[str, Any]]]:
    source, raw = _read_bytes(path, allowed_roots)
    extension = source.suffix.casefold()
    if extension in {".csv", ".tsv"}:
        text = raw.decode("utf-8-sig")
        delimiter = "\t" if extension == ".tsv" else ","
        if extension == ".csv" and text:
            try:
                delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
            except csv.Error:
                pass
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    elif extension == ".json":
        data = json.loads(raw.decode("utf-8-sig"))
        if isinstance(data, dict):
            rows = data.get("rows") if isinstance(data.get("rows"), list) else [data]
        else:
            rows = data
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("Le JSON doit contenir une liste d'objets ou un objet avec une liste rows.")
    elif extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("Lecture XLSX indisponible. Installe les dépendances du projet.") from error
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            workbook.close()
            raise ValueError("Le classeur XLSX ne contient aucune feuille active.")
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or f"colonne_{index + 1}") for index, value in enumerate(next(values, ()))]
        rows = [dict(zip(headers, row)) for row in values]
        workbook.close()
    else:
        raise ValueError("L'analyse de tableaux accepte uniquement CSV, TSV, JSON et XLSX.")
    if len(rows) > MAX_TABLE_ROWS:
        raise ValueError(f"Le tableau dépasse la limite de {MAX_TABLE_ROWS} lignes.")
    return source, rows


def _write_rows(path, rows, allowed_roots=None):
    target = _resolve(path, allowed_roots)
    extension = target.suffix.casefold()
    if extension == ".json":
        content = json.dumps(rows, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    elif extension in {".csv", ".tsv"}:
        if not rows:
            content = b""
        else:
            buffer = io.StringIO(newline="")
            writer = csv.DictWriter(buffer, fieldnames=list(rows[0]), delimiter="\t" if extension == ".tsv" else ",")
            writer.writeheader()
            writer.writerows(rows)
            content = buffer.getvalue().encode("utf-8")
    else:
        raise ValueError("Le résultat doit être enregistré en CSV, TSV ou JSON.")
    return _write_bytes(str(target), content, allowed_roots)


def _create_document(arguments, allowed_roots):
    path = _resolve(arguments.get("path"), allowed_roots)
    extension = (arguments.get("format") or path.suffix.lstrip(".")).casefold().lstrip(".")
    content = arguments.get("content", "")
    rows = arguments.get("rows")
    if extension in {"txt", "md", "markdown"}:
        if rows is not None:
            content = json.dumps(rows, ensure_ascii=False, indent=2)
        if not isinstance(content, str):
            raise ValueError("Le contenu du document doit être du texte.")
        return _write_bytes(str(path.with_suffix("." + extension)), content.encode("utf-8"), allowed_roots)
    if extension in {"csv", "tsv", "json"}:
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("Fournis rows sous forme de liste d'objets pour créer un tableau.")
        return _write_rows(str(path.with_suffix("." + extension)), rows, allowed_roots)
    if extension == "docx":
        try:
            from docx import Document
        except ImportError as error:
            raise ValueError("Création DOCX indisponible. Installe les dépendances du projet.") from error
        document = Document()
        for line in str(content).splitlines() or [""]:
            document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        return _write_bytes(str(path.with_suffix(".docx")), buffer.getvalue(), allowed_roots)
    if extension == "pdf":
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
        except ImportError as error:
            raise ValueError("Création PDF indisponible. Installe les dépendances du projet.") from error
        buffer = io.BytesIO()
        page = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter
        text = page.beginText(48, height - 54)
        text.setFont("Helvetica", 10)
        for line in str(content).splitlines() or [""]:
            safe_line = line.encode("cp1252", errors="replace").decode("cp1252")
            if text.getY() < 48:
                page.drawText(text)
                page.showPage()
                text = page.beginText(48, height - 54)
                text.setFont("Helvetica", 10)
            text.textLine(safe_line)
        page.drawText(text)
        page.save()
        return _write_bytes(str(path.with_suffix(".pdf")), buffer.getvalue(), allowed_roots)
    if extension == "xlsx":
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("Fournis rows sous forme de liste d'objets pour créer une feuille.")
        try:
            from openpyxl import Workbook
        except ImportError as error:
            raise ValueError("Création XLSX indisponible. Installe les dépendances du projet.") from error
        workbook = Workbook()
        sheet = workbook.active
        if sheet is None:
            raise ValueError("Impossible de créer la feuille active XLSX.")
        columns = list(rows[0]) if rows else []
        if columns:
            sheet.append(columns)
            for row in rows:
                sheet.append([row.get(column) for column in columns])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return _write_bytes(str(path.with_suffix(".xlsx")), buffer.getvalue(), allowed_roots)
    if extension == "pptx":
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError as error:
            raise ValueError("Création PPTX indisponible. Installe les dépendances du projet.") from error
        presentation = Presentation()
        for index, line in enumerate(str(content).splitlines() or [""]):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            if index == 0:
                title_shape = getattr(slide.shapes, "title", None)
                if title_shape is not None:
                    getattr(title_shape, "text_frame").text = line
            else:
                placeholder = slide.placeholders[1]
                getattr(placeholder, "text_frame").text = line
        buffer = io.BytesIO()
        presentation.save(buffer)
        return _write_bytes(str(path.with_suffix(".pptx")), buffer.getvalue(), allowed_roots)
    raise ValueError("Format pris en charge : TXT, Markdown, CSV, TSV, JSON, PDF, DOCX, XLSX ou PPTX.")


def _archive(arguments, allowed_roots):
    action = arguments.get("action")
    if action == "create_archive":
        destination = _resolve(arguments.get("destination"), allowed_roots)
        sources = arguments.get("paths")
        if not isinstance(sources, list) or not sources:
            raise ValueError("Fournis paths avec au moins un fichier ou dossier.")
        if not arguments.get("confirmed"):
            return {"confirmation_required": True, "message": "Confirme la création de l'archive ZIP demandée."}
        entries = []
        archive_names = set()
        total_size = 0
        for source_name in sources:
            source = _resolve(source_name, allowed_roots)
            if source.is_symlink() or not source.exists():
                raise ValueError("Les liens symboliques et chemins inexistants ne peuvent pas être archivés.")
            if source == destination or (source.is_dir() and source in destination.parents):
                raise ValueError("L'archive ne peut pas être créée dans un dossier qu'elle contient.")
            files = [source] if source.is_file() else [item for item in source.rglob("*") if item.is_file()]
            for item in files:
                if item.is_symlink():
                    continue
                archive_name = PurePosixPath(source.name, *item.relative_to(source).parts) if source.is_dir() else PurePosixPath(source.name)
                key = archive_name.as_posix().casefold()
                if key in archive_names:
                    raise ValueError("Les chemins fournis produisent des noms en double dans l'archive.")
                archive_names.add(key)
                total_size += item.stat().st_size
                if total_size > MAX_ARCHIVE_BYTES or len(entries) >= MAX_ARCHIVE_ENTRIES:
                    raise ValueError("Le contenu dépasse les limites de création d'archive.")
                entries.append((item, archive_name.as_posix()))
        destination.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(destination, "x", compression=zipfile.ZIP_DEFLATED) as archive:
            for item, archive_name in entries:
                archive.write(item, archive_name)
        return {"path": str(destination), "created": True, "size": destination.stat().st_size}

    if action == "extract_archive":
        source = _resolve(arguments.get("path"), allowed_roots)
        destination = _resolve(arguments.get("destination"), allowed_roots)
        if not arguments.get("confirmed"):
            return {"confirmation_required": True, "message": "Confirme l'extraction de l'archive ZIP dans le dossier demandé."}
        if not source.is_file() or source.stat().st_size > MAX_DOCUMENT_BYTES:
            raise ValueError("Archive absente ou trop volumineuse (maximum 20 Mo).")
        destination.mkdir(parents=True, exist_ok=True)
        root = destination.resolve()
        with zipfile.ZipFile(source) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES or sum(info.file_size for info in infos) > MAX_ARCHIVE_BYTES:
                raise ValueError("L'archive dépasse les limites d'extraction.")
            paths = set()
            for info in infos:
                name = info.filename
                normalized_name = name[:-1] if info.is_dir() and name.endswith("/") else name
                parts = normalized_name.split("/")
                if (
                    not normalized_name
                    or "\\" in name
                    or name.startswith("/")
                    or re.match(r"^[A-Za-z]:", name)
                    or any(part in {"", ".", ".."} for part in parts)
                ):
                    raise ValueError("L'archive contient un chemin dangereux.")
                if (info.external_attr >> 16) & 0o170000 == stat.S_IFLNK:
                    raise ValueError("L'extraction de liens symboliques est interdite.")
                target = (root / Path(*parts)).resolve()
                if root not in target.parents and target != root:
                    raise ValueError("L'archive tente d'écrire hors du dossier cible.")
                key = target.as_posix().casefold()
                if key in paths:
                    raise ValueError("L'archive contient des chemins en double.")
                paths.add(key)
                if target.exists():
                    raise FileExistsError(f"La cible existe déjà : {target}")
            archive.extractall(root)
        return {"path": str(root), "extracted": len(infos)}

    raise ValueError("Action d'archive inconnue.")


def _sqlite_query(path, query, allowed_roots):
    database = _resolve(path, allowed_roots)
    if not database.is_file() or database.stat().st_size > MAX_DOCUMENT_BYTES:
        raise ValueError("Base SQLite absente ou trop volumineuse.")
    if not isinstance(query, str) or not re.match(r"^\s*(SELECT|WITH)\b", query, re.I):
        raise ValueError("Seules les requêtes SELECT ou WITH sont permises.")
    if len(query) > 32000:
        raise ValueError("La requête dépasse la taille maximale de 32 Ko.")
    connection = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
    connection.execute("PRAGMA query_only = ON")
    denied = {
        sqlite3.SQLITE_INSERT, sqlite3.SQLITE_UPDATE, sqlite3.SQLITE_DELETE,
        sqlite3.SQLITE_CREATE_INDEX, sqlite3.SQLITE_CREATE_TABLE, sqlite3.SQLITE_CREATE_TRIGGER,
        sqlite3.SQLITE_CREATE_VIEW, sqlite3.SQLITE_DROP_INDEX, sqlite3.SQLITE_DROP_TABLE,
        sqlite3.SQLITE_DROP_TRIGGER, sqlite3.SQLITE_DROP_VIEW, sqlite3.SQLITE_ALTER_TABLE,
        sqlite3.SQLITE_ATTACH, sqlite3.SQLITE_DETACH, sqlite3.SQLITE_PRAGMA,
    }
    connection.set_authorizer(lambda action, *_args: sqlite3.SQLITE_DENY if action in denied else sqlite3.SQLITE_OK)
    connection.set_progress_handler(lambda: 1, 1_000_000)
    try:
        try:
            cursor = connection.execute(query)
        except sqlite3.OperationalError as error:
            if "interrupted" in str(error).casefold():
                raise ValueError("La requête a dépassé sa limite de traitement.") from error
            raise
        if cursor.description is None:
            raise ValueError("La requête n'a retourné aucun jeu de résultats.")
        columns = [item[0] for item in cursor.description]
        rows = cursor.fetchmany(MAX_SQL_ROWS + 1)
        return {"columns": columns, "rows": [dict(zip(columns, row)) for row in rows[:MAX_SQL_ROWS]], "truncated": len(rows) > MAX_SQL_ROWS}
    finally:
        connection.close()


def _chart(arguments, allowed_roots):
    path = _resolve(arguments.get("path"), allowed_roots)
    labels = arguments.get("labels")
    values = arguments.get("values")
    if not isinstance(labels, list) or not isinstance(values, list) or not labels or len(labels) != len(values) or len(labels) > 100:
        raise ValueError("Fournis 1 à 100 labels et valeurs de même longueur.")
    numbers = [float(value) for value in values]
    if any(not math.isfinite(value) for value in numbers):
        raise ValueError("Les valeurs du graphique doivent être finies.")
    width, height = 900, 520
    margin = 64
    maximum = max(max(numbers), 0.0) or 1.0
    minimum = min(min(numbers), 0.0)
    span = maximum - minimum or 1.0
    plot_height = height - margin * 2
    slot = (width - margin * 2) / len(numbers)
    zero_y = margin + maximum / span * plot_height
    bars = []
    for index, (label, value) in enumerate(zip(labels, numbers)):
        bar_height = abs(value) / span * plot_height
        y = zero_y - bar_height if value >= 0 else zero_y
        x = margin + index * slot + slot * 0.18
        bars.append(
            f'<rect x="{x:.2f}" y="{y:.2f}" width="{slot * 0.64:.2f}" height="{bar_height:.2f}" fill="#2f855a"/>'
            f'<text x="{x + slot * 0.32:.2f}" y="{height - 24}" text-anchor="middle">{html.escape(str(label)[:40])}</text>'
        )
    title = html.escape(str(arguments.get("title", "Graphique"))[:120])
    svg = (
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
        f'<rect width="100%" height="100%" fill="white"/><text x="{width / 2}" y="30" text-anchor="middle" font-size="20">{title}</text>'
        f'<line x1="{margin}" y1="{zero_y:.2f}" x2="{width-margin}" y2="{zero_y:.2f}" stroke="#777"/>'
        + "".join(bars) + "</svg>"
    )
    if path.suffix.casefold() != ".svg":
        path = path.with_suffix(".svg")
    return _write_bytes(str(path), svg.encode("utf-8"), allowed_roots)


def use(arguments, allowed_roots=None):
    if not isinstance(arguments, dict):
        raise ValueError("Les paramètres du document doivent être un objet.")
    action = arguments.get("action")
    if action in {"create_archive", "extract_archive"}:
        return _archive(arguments, allowed_roots)
    if action == "create_document":
        return _create_document(arguments, allowed_roots)
    if action == "read_document":
        path, raw = _read_bytes(arguments.get("path"), allowed_roots)
        content = _text_from_document(path, raw)
        if path.suffix.casefold() == ".pdf" and arguments.get("ocr"):
            content = content.strip() or _ocr_pdf(raw)
        return {"path": str(path), "format": path.suffix.lstrip("."), "content": content}
    if action == "convert_document":
        source, raw = _read_bytes(arguments.get("source"), allowed_roots)
        content = _text_from_document(source, raw)
        return _create_document({"path": arguments.get("destination"), "content": content}, allowed_roots)
    if action in {"analyze_table", "transform_table"}:
        source, rows = _rows_from_path(arguments.get("path"), allowed_roots)
        if action == "analyze_table":
            columns = list(rows[0]) if rows else []
            types = {
                column: sorted({type(row.get(column)).__name__ for row in rows if row.get(column) is not None})
                for column in columns
            }
            return {"path": str(source), "row_count": len(rows), "columns": columns, "types": types, "sample": rows[:10]}
        operation = arguments.get("operation")
        column = arguments.get("column")
        if operation == "sort":
            if not isinstance(column, str):
                raise ValueError("Le tri nécessite un nom de colonne.")

            def sort_key(row):
                value = row.get(column)
                try:
                    return (0, float(value))
                except (TypeError, ValueError):
                    return (1, str(value or "").casefold())

            rows.sort(key=sort_key, reverse=bool(arguments.get("descending")))
        elif operation == "filter":
            if not isinstance(column, str):
                raise ValueError("Le filtre nécessite un nom de colonne.")
            expected = str(arguments.get("value", "")).casefold()
            rows = [row for row in rows if expected in str(row.get(column, "")).casefold()]
        elif operation == "clean":
            cleaned = []
            seen = set()
            for row in rows:
                normalized = {key: value.strip() if isinstance(value, str) else value for key, value in row.items()}
                if not any(value not in (None, "") for value in normalized.values()):
                    continue
                fingerprint = json.dumps(normalized, sort_keys=True, ensure_ascii=False, default=str)
                if fingerprint not in seen:
                    cleaned.append(normalized)
                    seen.add(fingerprint)
            rows = cleaned
        else:
            raise ValueError("Opération prise en charge : sort, filter ou clean.")
        return _write_rows(arguments.get("destination"), rows, allowed_roots)
    if action == "sqlite_query":
        return _sqlite_query(arguments.get("path"), arguments.get("query"), allowed_roots)
    if action == "compare_metadata":
        left, _left_raw = _read_bytes(arguments.get("left"), allowed_roots)
        right, _right_raw = _read_bytes(arguments.get("right"), allowed_roots)
        left_stat, right_stat = left.stat(), right.stat()
        return {
            "left": str(left), "right": str(right),
            "same_size": left_stat.st_size == right_stat.st_size,
            "size_difference": right_stat.st_size - left_stat.st_size,
            "same_extension": left.suffix.casefold() == right.suffix.casefold(),
            "same_content": left.read_bytes() == right.read_bytes(),
        }
    if action == "create_chart":
        return _chart(arguments, allowed_roots)
    raise ValueError("Action document inconnue.")